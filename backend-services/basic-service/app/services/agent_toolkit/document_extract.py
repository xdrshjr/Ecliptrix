"""文档抽取：PDF / DOCX / 文本 / CSV / 字幕 / HTML → 分页纯文本。

设计基线：docs/plans/agent-general-toolkit/spec.md §3.3.4。

页口径按格式区分（R-11）：
- PDF 取**物理页**（`pypdf` 逐页 `extract_text()`）；
- 其他格式按 `DOCUMENT_PAGE_CHARS = 3000` 字符分块。
`document_max_pages` 默认 800 → 800 × 3000 = 2.4 M ≥ `document_max_chars` = 2 M，
两个上限不会互相否决。
"""
import asyncio
import hashlib
import io
from datetime import datetime, timedelta, timezone
from pathlib import Path

from app.models.agent_toolkit_models import AgentToolkitError

DOCUMENT_PAGE_CHARS = 3000
DIGEST_THRESHOLD_CHARS = 60_000
MAX_READ_CHARS = 8000
DEFAULT_READ_CHARS = 4000

# 扩展名 → kind（`AgentDocument.kind` 取值见 §5.1）
KIND_BY_EXTENSION = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "text",
    ".md": "text",
    ".csv": "csv",
    ".tsv": "csv",
    ".json": "text",
    ".srt": "subtitle",
    ".vtt": "subtitle",
    ".html": "html",
    ".htm": "html",
}
SUPPORTED_EXTENSIONS = tuple(sorted(KIND_BY_EXTENSION))
SUPPORTED_FORMATS = list(SUPPORTED_EXTENSIONS)


def utcnow():
    return datetime.now(timezone.utc)


def extension_of(filename):
    name = (filename or "").rsplit("/", 1)[-1]
    index = name.rfind(".")
    return name[index:].lower() if index > 0 else ""


def available():
    """`pypdf` / `python-docx` 任一缺失时整条文档能力降级（503）。"""
    try:
        import pypdf  # noqa: F401
        import docx  # noqa: F401
    except Exception:  # noqa: BLE001 - 可选依赖缺失不阻塞进程
        return False
    return True


def ensure_extension(filename):
    extension = extension_of(filename)
    if extension not in KIND_BY_EXTENSION:
        error = AgentToolkitError(
            "document_unsupported_format",
            message=f"unsupported document format: {extension or '(none)'}")
        error.supported = SUPPORTED_FORMATS
        raise error
    return extension, KIND_BY_EXTENSION[extension]


def _decode(data):
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise AgentToolkitError("document_parse_failed",
                            message="document is not decodable text")


def _extract_pdf(data):
    try:
        from pypdf import PdfReader
    except Exception:  # noqa: BLE001
        raise AgentToolkitError("documents_unavailable") from None
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            try:
                pages.append((page.extract_text() or "").strip())
            except Exception:  # noqa: BLE001 - 单页损坏不该毁掉整篇
                pages.append("")
    except AgentToolkitError:
        raise
    except Exception as error:  # noqa: BLE001
        # `reason` 只放异常类名，绝不回显文件内容（§3.3.4）。
        failure = AgentToolkitError(
            "document_parse_failed",
            message="pdf parse failed")
        failure.reason = type(error).__name__
        raise failure from None
    return pages


def _extract_docx(data):
    try:
        import docx
    except Exception:  # noqa: BLE001
        raise AgentToolkitError("documents_unavailable") from None
    try:
        document = docx.Document(io.BytesIO(data))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                blocks.append(" | ".join(cells))
    except AgentToolkitError:
        raise
    except Exception as error:  # noqa: BLE001
        failure = AgentToolkitError(
            "document_parse_failed",
            message="docx parse failed")
        failure.reason = type(error).__name__
        raise failure from None
    return "\n".join(block for block in blocks if block.strip())


def _extract_html(data):
    text = _decode(data)
    try:
        from bs4 import BeautifulSoup
    except Exception:  # noqa: BLE001
        return text
    soup = BeautifulSoup(text, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    return soup.get_text("\n")


def _normalize_subtitle(text):
    """SRT/VTT → 只留字幕正文，去掉序号与时间轴。"""
    lines = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line or line.isdigit() or "-->" in line:
            continue
        if line.upper().startswith("WEBVTT"):
            continue
        lines.append(line)
    return "\n".join(lines)


def extract(data, extension):
    """返回 `(text, source_pages)`；`source_pages` 只对 PDF 有意义。

    PDF 不再对拼接结果 `strip()`：`build_page_offsets` 按「每页正文 + 两个换行」
    推进游标，开头留白被 strip 掉会让**每一页**的偏移都漂移（首空白页时第 2 页
    只能读到半页）。
    """
    if extension == ".pdf":
        pages = _extract_pdf(data)
        return "\n\n".join(pages), len(pages)
    if extension == ".docx":
        return _extract_docx(data), None
    if extension in (".html", ".htm"):
        return _extract_html(data), None
    text = _decode(data)
    if extension in (".srt", ".vtt"):
        return _normalize_subtitle(text), None
    return text, None


def build_page_offsets(text, source_pages=None, physical_pages=None):
    """生成 `[{page, start, end}]`（字符偏移，页从 1 开始）。

    PDF 走物理页边界，其余按 3000 字符分块。`physical_pages` 是 PDF 每页的
    文本（用于把拼接后的全文按真实页边界切开）。

    `text` 可能已被 `truncate_text` 截断：超出的页必须**丢弃**而不是记成空页，
    否则 `read_document` 会把「不存在的页」报告成可读的空内容。
    """
    offsets = []
    if physical_pages is not None:
        cursor = 0
        for index, body in enumerate(physical_pages, start=1):
            start = cursor
            if start >= len(text):
                break
            end = min(start + len(body), len(text))
            offsets.append({"page": index, "start": start, "end": end})
            cursor = end + 2  # 拼接符 "\n\n"
        if offsets:
            return offsets
    for index in range(0, max(len(text), 1), DOCUMENT_PAGE_CHARS):
        start = index
        end = min(len(text), index + DOCUMENT_PAGE_CHARS)
        offsets.append({"page": len(offsets) + 1, "start": start, "end": end})
    return offsets or [{"page": 1, "start": 0, "end": 0}]


def split_pdf_pages(data):
    """PDF 单独返回每页文本，供真实物理页偏移使用。"""
    return _extract_pdf(data)


def page_range(offsets, page):
    if not isinstance(page, int) or page < 1 or page > len(offsets):
        error = AgentToolkitError(
            "page_out_of_range",
            message=f"page must be 1..{len(offsets)}")
        error.pages = len(offsets)
        raise error
    entry = offsets[page - 1]
    return int(entry["start"]), int(entry["end"])


def slice_page(text, offsets, page, offset=0, max_chars=DEFAULT_READ_CHARS):
    """按页切片；单页超长时返回 `next_offset`（页内字符续读，R-10）。"""
    start, end = page_range(offsets, page)
    if not 0 <= offset <= (end - start):
        raise AgentToolkitError("invalid_offset")
    if not 1 <= max_chars <= MAX_READ_CHARS:
        raise AgentToolkitError("invalid_page_size",
                                message=f"max_chars must be 1..{MAX_READ_CHARS}")
    body = text[start:end]
    chunk = body[offset: offset + max_chars]
    consumed = offset + len(chunk)
    return {
        "content": chunk,
        "next_offset": consumed if consumed < len(body) else None,
        "truncated": consumed < len(body),
    }


def truncate_text(text, max_chars):
    """超出 `document_max_chars` 截断并报告丢弃量（不报错）。"""
    if len(text) <= max_chars:
        return text, False, 0
    return text[:max_chars], True, len(text) - max_chars


def sha256_bytes(data):
    return hashlib.sha256(data).hexdigest()


async def write_extracted(directory, document_id, text):
    """`extracted/<document_id>.txt`（UTF-8）。"""
    target = Path(directory) / "extracted" / (document_id + ".txt")

    def _write():
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(text, encoding="utf-8")

    await asyncio.to_thread(_write)
    return target


async def write_upload(directory, document_id, extension, data):
    """`uploads/<document_id>.<ext>`：原件保真，供重解析/审计。"""
    target = Path(directory) / "uploads" / (document_id + extension)

    def _write():
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = target.with_name(target.name + ".tmp")
        temporary.write_bytes(data)
        temporary.replace(target)

    await asyncio.to_thread(_write)
    return target


def retention_deadline(days):
    if days is None or int(days) <= 0:
        return None
    return utcnow().replace(tzinfo=None) + timedelta(days=int(days))


# ── 摄取 / 列表 / 读取 / 删除 ──────────────────────────────────────────────


def _session_factory(factory=None):
    if factory is not None:
        return factory
    from app.database import async_session_factory

    return async_session_factory


def _settings():
    from app.services.system_settings_service import toolkit_settings

    return toolkit_settings()


def _gate():
    settings = _settings()
    if not settings["enabled"]:
        raise AgentToolkitError("toolkit_disabled")
    if not settings["documents_enabled"]:
        raise AgentToolkitError("documents_disabled")
    if not available():
        raise AgentToolkitError(
            "documents_unavailable",
            message="pypdf / python-docx are not installed")
    return settings


async def ingest(filename, data, project_id, session_factory=None, root=None):
    """上传 → 抽取 → 落盘 → 入库 → 触发摘要（不阻塞响应）。"""
    from uuid import uuid4

    from app.services.agent_toolkit import document_digest, workspace_store
    from app.services.agent_toolkit.workspace_store import project_directory

    settings = _gate()
    if len(data) == 0:
        raise AgentToolkitError("invalid_argument", message="empty upload")
    if len(data) > workspace_store.max_file_bytes():
        error = AgentToolkitError(
            "document_too_large",
            message=f"file exceeds {int(settings['max_file_mb'])} MiB")
        error.limit = int(settings["max_file_mb"])
        raise error
    extension, kind = ensure_extension(filename)
    if extension == ".pdf":
        # 一次解析：正文与页偏移必须来自同一份页列表，否则两处口径迟早漂移。
        physical = split_pdf_pages(data)
        text = "\n\n".join(physical)
        source_pages = len(physical)
    else:
        physical = None
        text, source_pages = extract(data, extension)
    text, truncated, dropped = truncate_text(
        text, int(settings["document_max_chars"]))
    offsets = build_page_offsets(text, physical_pages=physical)
    if len(offsets) > int(settings["document_max_pages"]):
        error = AgentToolkitError(
            "document_too_large",
            message=f"document has more than "
                    f"{int(settings['document_max_pages'])} pages")
        error.pages = len(offsets)
        error.limit = int(settings["document_max_pages"])
        raise error
    document_id = str(uuid4())
    directory = workspace_store.ensure_layout(project_directory(project_id, root))
    await write_upload(directory, document_id, extension, data)
    await write_extracted(directory, document_id, text)
    relative_text = (
        Path(directory).name + "/extracted/" + document_id + ".txt")
    now = utcnow().replace(tzinfo=None)
    display_name = (filename or "document").rsplit("/", 1)[-1][:200]
    row = _document_row(
        id=document_id, project_id=project_id, name=display_name,
        kind=kind, pages=len(offsets), chars=len(text),
        text_path=relative_text, page_offsets=offsets,
        sha256=sha256_bytes(data), created_at=now,
        digest_state="running" if document_digest.should_digest(text) else "skipped",
    )
    factory = _session_factory(session_factory)
    async with factory() as session:
        session.add(row)
        await session.commit()
    document_digest.schedule(document_id, text, offsets, directory, factory)
    return {
        "document_id": document_id,
        "name": display_name,
        "kind": kind,
        "pages": len(offsets),
        "chars": len(text),
        "source_pages": source_pages if source_pages is not None else len(offsets),
        "truncated": truncated,
        "dropped_chars": dropped,
        "digest_status": row.digest_state,
    }


def _document_row(**values):
    from app.models.agent_toolkit_db_models import AgentDocument

    return AgentDocument(**values)


async def list_documents(project_id, limit=50, session_factory=None):
    from sqlalchemy import select

    from app.models.agent_toolkit_db_models import AgentDocument
    from app.models.agent_toolkit_models import normalize_project_id

    _gate()
    project = normalize_project_id(project_id)
    factory = _session_factory(session_factory)
    async with factory() as session:
        rows = list((await session.execute(select(AgentDocument).where(
            AgentDocument.state == "ready"))).scalars())
    rows = [row for row in rows if (row.project_id or None) == project]
    rows.sort(key=lambda row: row.created_at, reverse=True)
    return {
        "documents": [{
            "id": str(row.id),
            "document_id": str(row.id),
            "name": row.name,
            "kind": row.kind,
            "pages": row.pages,
            "chars": row.chars,
            "created_at": row.created_at.isoformat(),
        } for row in rows[: int(limit)]],
        "total": len(rows),
    }


async def _load(document_id, project_id, session_factory=None):
    from sqlalchemy import select

    from app.models.agent_toolkit_db_models import AgentDocument
    from app.models.agent_toolkit_models import normalize_project_id

    project = normalize_project_id(project_id)
    factory = _session_factory(session_factory)
    async with factory() as session:
        row = (await session.execute(select(AgentDocument).where(
            AgentDocument.id == str(document_id)))).scalar_one_or_none()
    if row is None or row.state != "ready" or (row.project_id or None) != project:
        raise AgentToolkitError("document_not_found")
    return row


async def read_document(project_id, document_id=None, name=None, page=1,
                        offset=0, max_chars=DEFAULT_READ_CHARS,
                        session_factory=None, root=None):
    """按页读取；**永不**整篇返回（AC-11）。"""
    from app.services.agent_toolkit import document_digest, workspace_store
    from app.services.agent_toolkit.workspace_store import project_directory

    _gate()
    if document_id is None:
        if not name:
            raise AgentToolkitError(
                "invalid_argument",
                message="pass document_id or name")
        found = await _find_by_name(project_id, name, session_factory)
        document_id = found
    row = await _load(document_id, project_id, session_factory)
    root_path = Path(root) if root is not None else workspace_store.workspace_root()
    directory = project_directory(project_id, root)
    text_file = root_path / row.text_path
    if not text_file.is_file():
        raise AgentToolkitError("document_not_found")
    text = text_file.read_text(encoding="utf-8")
    sliced = slice_page(text, row.page_offsets, int(page),
                        int(offset), int(max_chars))
    digest = document_digest.read_digest(directory, str(row.id))
    payload = {
        "document_id": str(row.id),
        "name": row.name,
        "kind": row.kind,
        "page": int(page),
        "pages": len(row.page_offsets),
        "source_pages": len(row.page_offsets),
        "next_page": int(page) + 1 if int(page) < len(row.page_offsets) else None,
        "digest_status": row.digest_state,
        "digest_is_summary": True,
        "dropped_chars": 0,
        **sliced,
    }
    if digest:
        payload["digest"] = digest.get("digest")
    return payload


async def _find_by_name(project_id, name, session_factory=None):
    listing = await list_documents(project_id, 100, session_factory)
    for entry in listing["documents"]:
        if entry["name"] == name:
            return entry["id"]
    raise AgentToolkitError("document_not_found",
                            message=f"no document named {name!r}")


async def delete_document(project_id, document_id, session_factory=None,
                          root=None):
    from app.services.agent_toolkit.workspace_store import project_directory

    _gate()
    from app.services.agent_toolkit import workspace_store

    row = await _load(document_id, project_id, session_factory)
    root_path = Path(root) if root is not None else workspace_store.workspace_root()
    directory = project_directory(project_id, root)
    row.state = "deleted"
    factory = _session_factory(session_factory)
    async with factory() as session:
        stored = await session.get(type(row), row.id)
        if stored is not None:
            stored.state = "deleted"
            await session.commit()
    text_file = root_path / row.text_path
    text_file.unlink(missing_ok=True)
    (Path(directory) / "digest" / (str(row.id) + ".json")).unlink(missing_ok=True)
    for path in (Path(directory) / "uploads").glob(str(row.id) + ".*"):
        path.unlink(missing_ok=True)
    return {"deleted": True}

